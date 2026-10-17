"""复盘录音逐字稿文档：无时间码、无编号，按课程/小节分段。

用法：
    python review_doc.py --keep day2_keep.json --asr day2_asr.json \
        --sections sections.json --title "9月10日 学习复盘" \
        --meta "2026年9月10日（周四）" -o day2_transcript

sections.json 格式（**必须覆盖 keep 里的每一句**，脚本会断言检查）：
    [["开场", [0, 1]], ["美术课", [3, 4, 5]], ...]

产出：<out>.md / .txt / .html / .docx（docx 需 python-docx，缺失会自动跳过）
docx 中文字体要额外设 w:eastAsia，字体从 profile.json 读取。
"""
import argparse
import json
from pathlib import Path

from config import load_profile

PROFILE = load_profile()


def end_punct(t):
    t = t.strip()
    if not t:
        return t
    # 以「是」结尾说明话没说完，后面还要接内容，不能断句
    if t[-1] in "是":
        return t
    return t if t[-1] in "。！？，、；：.!?" else t + "。"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--keep", required=True)
    ap.add_argument("--asr", required=True)
    ap.add_argument("--sections", required=True)
    ap.add_argument("--title", default="学习复盘")
    ap.add_argument("--meta", default="")
    ap.add_argument("-o", "--out", default="transcript")
    ap.add_argument("--fix", default=None,
                    help="成稿替换表 json：[[错, 对], ...]，只影响文档不影响音频")
    a = ap.parse_args()

    keep = json.loads(Path(a.keep).read_text(encoding="utf-8-sig"))
    asr = json.loads(Path(a.asr).read_text(encoding="utf-8-sig"))
    CAP = {c["idx"]: c for c in keep["captions"]}
    KEEP = keep["keep"]
    sections = json.loads(Path(a.sections).read_text(encoding="utf-8-sig"))
    FIX = json.loads(Path(a.fix).read_text(encoding="utf-8-sig")) if a.fix else []

    covered = {x for _, xs in sections for x in xs}
    missing = [i for i in KEEP if i not in covered]
    if missing:
        raise SystemExit(f"有句子未归入任何小节: {missing}")

    seg_total = sum(b - s for s, b in keep["segments"]) / 1000
    orig = asr["duration_ms"] / 1000
    pct = round(100 - seg_total / orig * 100)

    def mmss(x):
        return f"{int(x//60)}:{x%60:05.2f}"

    def body(idxs):
        s = "".join(end_punct(CAP[i]["text"].strip()) for i in idxs)
        s = s.replace("。。", "。").replace("？。", "？").replace("！。", "！")
        for old, new in FIX:
            s = s.replace(old, new)
        return s

    out = Path(a.out)

    md = [f"# {a.title}", ""]
    if a.meta:
        md.append(f"> {a.meta}｜成片 {mmss(seg_total)}（原始 {mmss(orig)}，压缩 {pct}%）")
        md.append("")
    for name, idxs in sections:
        md += [f"## {name}", "", body(idxs), ""]
    out.with_suffix(".md").write_text("\n".join(md), encoding="utf-8")

    txt = [a.title]
    if a.meta:
        txt.append(f"{a.meta}｜成片 {mmss(seg_total)}｜原始 {mmss(orig)}｜压缩 {pct}%")
    txt.append("")
    for name, idxs in sections:
        txt += [name, body(idxs), ""]
    out.with_suffix(".txt").write_text("\n".join(txt), encoding="utf-8")

    blocks = "".join(f'<section><h2>{n}</h2><p>{body(x)}</p></section>'
                     for n, x in sections)
    html = f"""<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>{a.title}</title>
<style>*{{box-sizing:border-box}}body{{margin:0;background:#f5f5f7;color:#1d2129;
font:16px/1.9 "Songti SC","Noto Serif SC",Georgia,serif}}
.page{{max-width:760px;margin:0 auto;padding:48px 32px 80px;background:#fff;
box-shadow:0 1px 3px rgba(0,0,0,.08)}}h1{{font-size:26px;margin:0 0 6px}}
.meta{{color:#8a8f99;font-size:13px;margin-bottom:32px;font-family:-apple-system,"PingFang SC",sans-serif}}
h2{{font-size:15px;margin:34px 0 10px;color:#d91ad9;letter-spacing:.08em;
font-family:-apple-system,"PingFang SC",sans-serif}}
h2::before{{content:"";display:inline-block;width:3px;height:13px;background:#d91ad9;
margin-right:8px;vertical-align:-1px}}p{{margin:0 0 18px;text-align:justify}}
footer{{margin-top:48px;padding-top:16px;border-top:1px solid #ececf0;color:#a9aeb8;font-size:12px}}
@media print{{body{{background:#fff}}.page{{box-shadow:none;padding:0}}}}</style></head>
<body><div class="page"><h1>{a.title}</h1>
<div class="meta">{a.meta}｜成片 {mmss(seg_total)}｜原始 {mmss(orig)}｜压缩 {pct}%</div>
{blocks}
<footer>由录音转写并精剪生成 · 已去除语气词、重复与无关生活内容</footer>
</div></body></html>"""
    out.with_suffix(".html").write_text(html, encoding="utf-8")

    total_chars = sum(len(body(x)) for _, x in sections)
    print(f"小节 {len(sections)}，正文 {total_chars} 字")
    for name, idxs in sections:
        print(f"  {name}: {len(body(idxs))} 字 / {len(idxs)} 句")

    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn

        doc = Document()
        st = doc.styles["Normal"]
        font = PROFILE["editing"].get("docx_font", "微软雅黑")
        body_pt = int(PROFILE["editing"].get("docx_body_pt", 11))
        st.font.name = font
        st.font.size = Pt(body_pt)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), font)   # 不设这行中文不会生效
        h = doc.add_heading(a.title, level=0)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{a.meta}｜成片 {mmss(seg_total)}｜原始 {mmss(orig)}｜压缩 {pct}%")
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x8A, 0x8F, 0x99)
        for name, idxs in sections:
            hh = doc.add_heading(name, level=2)
            for run in hh.runs:
                run.font.color.rgb = RGBColor(0xD9, 0x1A, 0xD9)
            doc.add_paragraph(body(idxs)).paragraph_format.space_after = Pt(10)
        fp = doc.add_paragraph()
        fr = fp.add_run("由录音转写并精剪生成 · 已去除语气词、重复与无关生活内容")
        fr.font.size = Pt(8)
        fr.font.color.rgb = RGBColor(0xA9, 0xAE, 0xB8)
        doc.save(out.with_suffix(".docx"))
        print("docx:", out.with_suffix(".docx").name)
    except Exception as e:
        print("docx 跳过:", e)


if __name__ == "__main__":
    main()
