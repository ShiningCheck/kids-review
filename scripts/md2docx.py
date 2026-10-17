# -*- coding: utf-8 -*-
"""逐字稿.md -> 逐字稿.docx（同名）

    python md2docx.py [逐字稿_2026年09月14日.md]
    不带参数则取当前目录最新的 逐字稿*.md
"""
import re, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from config import load_profile

PROFILE = load_profile()
FONT = PROFILE["editing"].get("docx_font", "微软雅黑")
BODY_PT = int(PROFILE["editing"].get("docx_body_pt", 11))
src = sys.argv[1] if len(sys.argv) > 1 else None
if not src:
    cands = sorted(Path(".").glob("逐字稿*.md"), key=lambda p: p.stat().st_mtime)
    if not cands:
        raise SystemExit("没找到 逐字稿*.md")
    src = str(cands[-1])
docx_out = str(Path(src).with_suffix(".docx"))

doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(BODY_PT)
st.element.rPr.rFonts.set(
    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", FONT)
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.2)

lines = Path(src).read_text(encoding="utf-8-sig").splitlines()


def para(text="", size=BODY_PT, bold=False, color=None, style=None, space=4):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space)
    return p


for ln in lines:
    if not ln.strip():
        continue
    if ln.startswith("## "):
        para(ln[3:], size=14, bold=True, color=RGBColor(0x1F, 0x4E, 0x79), space=8)
        continue
    if ln.startswith("# "):
        p = para(ln[2:], size=19, bold=True, space=6)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        continue
    if ln.startswith("> "):
        para(ln[2:], size=9, color=RGBColor(0x77, 0x77, 0x77), space=2)
        continue
    if ln.startswith("---"):
        continue
    if ln.startswith("- `"):
        m = re.match(r"- `(.*?)` (.*)", ln)
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(f"{m.group(1)}  {m.group(2)}")
        r.font.name = FONT
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        p.paragraph_format.space_after = Pt(0)
        continue
    m = re.match(r"\*\*\[(.*?)\] (.*?)：\*\*(.*)", ln)
    if m:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r0 = p.add_run(f"[{m.group(1)}] ")
        r0.font.name, r0.font.size = FONT, Pt(9)
        r0.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
        r1 = p.add_run(f"{m.group(2)}：")
        r1.font.name, r1.font.size, r1.bold = FONT, Pt(BODY_PT), True
        r1.font.color.rgb = (RGBColor(0xC0, 0x6A, 0x2B) if m.group(2) == "大人"
                             else RGBColor(0x2E, 0x74, 0xB5))
        r2 = p.add_run(m.group(3))
        r2.font.name, r2.font.size = FONT, Pt(BODY_PT)
    else:
        para(ln, size=BODY_PT, space=3)

doc.save(docx_out)
print(docx_out)
